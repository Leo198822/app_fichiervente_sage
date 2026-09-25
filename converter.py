"""Conversion des exports Sage (écritures + clients) au format d'import Pennylane."""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font

import config

COLONNES_PENNYLANE = [
    "Date",
    "Code Journal",
    "Numéro de compte",
    "Libellé de compte",
    "Libellé de ligne",
    "Taux de TVA du compte",
    "Code pays du compte",
    "Libellé de pièce",
    "Numéro de pièce",
    "Débit et/ou Crédit",
    "Crédit",
    "Famille de catégories",
    "Catégorie",
    "Identifiant de ligne",
    "Identifiant de lettrage",
]

# Positions (début, fin) des champs dans les exports Sage à largeur fixe.
CHAMPS_CLIENT = {
    "code": (0, 13),
    "nom": (13, 37),
    "collectif": (50, 63),
    "adresse1": (63, 87),
    "adresse2": (87, 111),
    "code_postal": (111, 119),
    "ville": (119, 189),
    "pays": (189, 192),
}

CHAMPS_ECRITURE = {
    "journal": (0, 3),
    "date": (3, 9),
    "type_piece": (9, 11),
    "compte": (11, 24),
    "type_compte": (24, 25),
    "code_client": (25, 38),
    "reference": (38, 51),
    "libelle": (51, 76),
    "reglement": (76, 77),
    "echeance": (77, 83),
    "sens": (83, 84),
    "montant": (84, 104),
    "numero": (105, 118),
}

RE_FACTURE = re.compile(r"\b(F[AR]\d{5,})\b")
RE_ECHAPPEMENT = re.compile(r"\\(.)")


@dataclass
class Resultat:
    ecritures: pd.DataFrame
    alertes: list[str] = field(default_factory=list)
    nb_pieces: int = 0


def _decoder(contenu: bytes | str) -> list[str]:
    if isinstance(contenu, bytes):
        contenu = contenu.decode(config.ENCODAGE_SAGE)
    lignes = contenu.splitlines()
    # 1re ligne = nom de la société ; on ignore les lignes vides
    return [l for l in lignes[1:] if l.strip()]


def _extraire(ligne: str, champs: dict[str, tuple[int, int]]) -> dict[str, str]:
    # Sage échappe certains caractères (ex. "\#") ce qui décale les colonnes
    ligne = RE_ECHAPPEMENT.sub(r"\1", ligne)
    return {nom: ligne[d:f].strip() for nom, (d, f) in champs.items()}


def lire_clients(contenu: bytes | str) -> dict[str, dict[str, str]]:
    clients = {}
    for ligne in _decoder(contenu):
        client = _extraire(ligne, CHAMPS_CLIENT)
        if client["code"]:
            clients[client["code"]] = client
    return clients


def lire_ecritures(contenu: bytes | str) -> list[dict[str, str]]:
    return [_extraire(ligne, CHAMPS_ECRITURE) for ligne in _decoder(contenu)]


def _par_prefixe(valeur: str, table: dict[str, str]) -> str | None:
    for prefixe in sorted(table, key=len, reverse=True):
        if valeur.startswith(prefixe):
            return table[prefixe]
    return None


def convertir_journal(journal: str) -> str:
    return _par_prefixe(journal, config.JOURNAUX) or journal


def convertir_compte(compte: str) -> str:
    remplace = _par_prefixe(compte, config.COMPTES_REMPLACES)
    if remplace:
        return remplace
    if config.LONGUEUR_COMPTE and compte.isdigit():
        return compte.ljust(config.LONGUEUR_COMPTE, "0")
    return compte


def _date(ddmmyy: str) -> date:
    return datetime.strptime(ddmmyy, "%d%m%y").date()


def _montant(texte: str) -> float:
    return round(float(texte.replace(",", ".") or 0), 2)


def _signe(e: dict[str, str]) -> float:
    """Montant signé côté crédit (produits et TVA collectée positifs)."""
    return _montant(e["montant"]) * (1 if e["sens"] == "C" else -1)


def taux_tva(piece: list[dict[str, str]], comptes_tva: tuple[str, ...] = config.COMPTES_TVA_COLLECTEE) -> tuple[str, bool]:
    """Taux de TVA de la pièce (ex. "20 %", "pas de TVA") et indicateur de taux incohérent."""
    tva = sum(_signe(e) for e in piece if e["compte"].startswith(comptes_tva))
    ht = sum(_signe(e) for e in piece if e["compte"].startswith(config.COMPTES_SOUMIS_TVA))
    if abs(tva) < 0.005:
        return config.LIBELLE_SANS_TVA, False
    if abs(ht) < 0.005:
        return "", True
    calcule = tva / ht * 100
    taux = min(config.TAUX_TVA, key=lambda t: abs(t - calcule))
    # espace insécable avant "%", comme dans la trame Pennylane ("20 %")
    return f"{taux:g}\u00a0%".replace(".", ","), abs(taux - calcule) > 0.5


def _decouper_pieces(ecritures: list[dict[str, str]]) -> list[list[dict[str, str]]]:
    """Regroupe les lignes consécutives d'un même journal jusqu'à équilibre débit = crédit."""
    pieces, courante, solde = [], [], 0.0
    for e in ecritures:
        if courante and e["journal"] != courante[0]["journal"]:
            pieces.append(courante)
            courante, solde = [], 0.0
        courante.append(e)
        solde += _montant(e["montant"]) * (1 if e["sens"] == "D" else -1)
        if abs(solde) < 0.005:
            pieces.append(courante)
            courante, solde = [], 0.0
    if courante:
        pieces.append(courante)
    return pieces


def convertir(contenu_ecritures: bytes | str, contenu_clients: bytes | str) -> Resultat:
    clients = lire_clients(contenu_clients)
    ecritures = lire_ecritures(contenu_ecritures)
    alertes: list[str] = []
    clients_inconnus: set[str] = set()
    compteurs: dict[str, int] = {}
    lignes = []

    pieces = _decouper_pieces(ecritures)
    for piece in pieces:
        journal = convertir_journal(piece[0]["journal"])
        debit = sum(_montant(e["montant"]) for e in piece if e["sens"] == "D")
        credit = sum(_montant(e["montant"]) for e in piece if e["sens"] == "C")

        facture = next((m.group(1) for e in piece if (m := RE_FACTURE.search(e["libelle"]))), None)
        if facture:
            numero = facture
        else:
            d = _date(piece[0]["date"])
            cle = f"{journal}-{d:%y%m}"
            compteurs[cle] = compteurs.get(cle, 0) + 1
            numero = f"{cle}-{compteurs[cle]:03d}"

        if abs(debit - credit) >= 0.005:
            alertes.append(
                f"Pièce {numero} ({piece[0]['journal']} du {_date(piece[0]['date']):%d/%m/%Y}) "
                f"déséquilibrée : débit {debit:.2f} / crédit {credit:.2f}"
            )

        taux, taux_incoherent = taux_tva(piece)
        if taux_incoherent:
            alertes.append(f"Pièce {numero} : taux de TVA incohérent ({taux or 'HT nul'}), à vérifier")

        # Nom du client de la pièce : sert de libellé de ligne et de pièce pour toutes ses lignes
        code_piece = next((e["code_client"] for e in piece if e["type_compte"] == "X" and e["code_client"]), None)
        nom_piece = clients[code_piece]["nom"] if code_piece in clients else code_piece

        for e in piece:
            code_client = e["code_client"]
            if e["type_compte"] == "X" and code_client:
                compte = (config.PREFIXE_COMPTE_CLIENT + code_client).ljust(config.LONGUEUR_COMPTE or 0, "0")
                if code_client in clients:
                    libelle_compte = clients[code_client]["nom"]
                else:
                    libelle_compte = code_client
                    clients_inconnus.add(code_client)
            else:
                compte = convertir_compte(e["compte"])
                libelle_compte = _par_prefixe(compte, config.LIBELLES_COMPTES) or ""
            libelle = nom_piece or e["libelle"]

            montant = _montant(e["montant"])
            lignes.append({
                "Date": _date(e["date"]),
                "Code Journal": journal,
                "Numéro de compte": compte,
                "Libellé de compte": libelle_compte,
                "Libellé de ligne": libelle,
                "Taux de TVA du compte": taux if e["compte"].startswith(config.COMPTES_SOUMIS_TVA) else "",
                "Code pays du compte": "",
                "Libellé de pièce": libelle,
                "Numéro de pièce": numero,
                "Débit et/ou Crédit": montant if e["sens"] == "D" else 0.0,
                "Crédit": montant if e["sens"] == "C" else 0.0,
                "Famille de catégories": "",
                "Catégorie": "",
                "Identifiant de ligne": "",
                "Identifiant de lettrage": "",
            })

    for code in sorted(clients_inconnus):
        alertes.append(f"Code client {code} absent du fichier clients (le code est utilisé comme libellé)")

    df = pd.DataFrame(lignes, columns=COLONNES_PENNYLANE)
    return Resultat(ecritures=df, alertes=alertes, nb_pieces=len(pieces))


def vers_excel(df: pd.DataFrame) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Ecritures"
    ws.append(COLONNES_PENNYLANE)
    for cellule in ws[1]:
        cellule.font = Font(bold=True)
    for ligne in df.itertuples(index=False):
        ws.append([None if v == "" else v for v in ligne])
    for (col, fmt) in (("A", "yyyy-mm-dd"), ("J", "0.00"), ("K", "0.00")):
        for cellule in ws[col][1:]:
            cellule.number_format = fmt
    for col in ("C", "I"):
        for cellule in ws[col][1:]:
            cellule.number_format = "@"
    largeurs = {"A": 12, "B": 12, "C": 14, "D": 28, "E": 30, "H": 30, "I": 16, "J": 16, "K": 12}
    for col, largeur in largeurs.items():
        ws.column_dimensions[col].width = largeur
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
