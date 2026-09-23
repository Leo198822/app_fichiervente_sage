import io

from openpyxl import load_workbook

from converter import COLONNES_PENNYLANE, convertir, taux_tva, vers_excel


def ligne_client(code, nom, pays):
    return f"{code:<13}{nom:<24}{code:<13}{'41100000':<13}{'1 rue X':<48}{'75001':<8}{'Paris':<70}{pays:<3}"


def ligne_ecriture(journal, date, compte, sens, montant, libelle, client="", ref=""):
    aux = f"X{client:<13}" if client else " " * 14
    return (
        # Sage ajoute un "\" d'échappement en plus de la largeur du champ
        f"{journal:<3}{date}FC{compte:<13}{aux}{ref:<{13 + ref.count(chr(92))}}{libelle:<25}C{date}{sens}"
        f"{montant:>20}N{'1':<13}"
    )


CLIENTS = "\r\n".join([
    "BM Textile",
    ligne_client("3184", "Anne-Laure Taillefer", "FRA"),
    ligne_client("3103", "Timea Griset", "USA"),
]).encode("cp850")

ECRITURES = "\r\n".join([
    "BM Textile",
    ligne_ecriture("VTE", "020726", "41100000", "D", "252.90", "Facture n° FA2600304 Anne", "3184", "PO\\#1"),
    ligne_ecriture("VTE", "020726", "44571120", "C", "42.15", "Facture n° FA2600304 Anne"),
    ligne_ecriture("VTE", "020726", "70702000", "C", "210.75", "Facture n° FA2600304 Anne"),
    ligne_ecriture("BQ1", "030726", "41100000", "C", "252.90", "0555211", "3184"),
    ligne_ecriture("BQ1", "030726", "51211000", "D", "252.90", "0555211"),
    ligne_ecriture("OD ", "030726", "41100000", "D", "12.58", "Ecart", "9999"),
    ligne_ecriture("OD ", "030726", "76600000", "C", "12.58", "Ecart"),
    ligne_ecriture("CAI", "030726", "41100000", "D", "10.01", "remboursement", "3103"),
    ligne_ecriture("CAI", "310726", "53110000", "C", "10.01", "Centralisation"),
    "",
]).encode("cp850")


def test_conversion():
    r = convertir(ECRITURES, CLIENTS)
    df = r.ecritures
    assert list(df.columns) == COLONNES_PENNYLANE
    assert len(df) == 9
    assert r.nb_pieces == 4

    vente = df.iloc[0]
    assert vente["Code Journal"] == "VT"
    assert vente["Numéro de compte"] == "411003184"
    assert vente["Libellé de compte"] == "Anne-Laure Taillefer"
    assert vente["Code pays du compte"] == ""
    assert vente["Numéro de pièce"] == "FA2600304"
    # toutes les lignes de la pièce portent le nom du client
    assert set(df.iloc[0:3]["Libellé de ligne"]) == {"Anne-Laure Taillefer"}
    assert set(df.iloc[0:3]["Libellé de pièce"]) == {"Anne-Laure Taillefer"}
    assert vente["Débit et/ou Crédit"] == 252.90 and vente["Crédit"] == 0
    assert df.iloc[1]["Numéro de compte"] == "445711200"
    assert list(df.iloc[0:3]["Taux de TVA du compte"]) == ["", "", "20%"]
    assert df.iloc[6]["Taux de TVA du compte"] == "pas de TVA"  # 766 sans TVA

    banque = df.iloc[4]
    assert banque["Code Journal"] == "SAGE"
    assert banque["Numéro de compte"] == "582000000"
    assert banque["Taux de TVA du compte"] == ""
    assert df.iloc[3]["Numéro de pièce"] == banque["Numéro de pièce"] == "SAGE-2607-001"

    assert df.iloc[5]["Code Journal"] == "SAGE"
    assert df.iloc[5]["Numéro de pièce"] == "SAGE-2607-002"
    assert df.iloc[6]["Libellé de ligne"] == "9999"
    assert r.alertes == ["Code client 9999 absent du fichier clients (le code est utilisé comme libellé)"]
    caisse = df.iloc[8]
    assert (caisse["Code Journal"], caisse["Numéro de compte"]) == ("SAGE", "531")
    assert caisse["Libellé de ligne"] == "Timea Griset"
    assert df["Débit et/ou Crédit"].sum() == df["Crédit"].sum()


def test_taux_tva():
    def piece(*lignes):
        return [{"compte": c, "sens": sens, "montant": m} for c, sens, m in lignes]

    assert taux_tva(piece(("44571120", "C", "10.00"), ("70702000", "C", "100.00"))) == ("10%", False)
    assert taux_tva(piece(("44571120", "C", "5.50"), ("70702000", "C", "100.00"))) == ("5,5%", False)
    # avoir : TVA et HT au débit
    assert taux_tva(piece(("44571120", "D", "1.67"), ("70702000", "D", "8.34"))) == ("20%", False)
    assert taux_tva(piece(("70709000", "C", "100.00"))) == ("pas de TVA", False)
    # mélange de lignes taxées et exonérées -> taux incohérent signalé
    assert taux_tva(piece(("44571120", "C", "20.00"), ("70702000", "C", "100.00"), ("70709000", "C", "50.00")))[1]


def test_piece_desequilibree():
    ecritures = "\r\n".join([
        "Soc",
        ligne_ecriture("VTE", "020726", "41100000", "D", "100.00", "Facture n° FA2600001 X", "3184"),
        ligne_ecriture("VTE", "020726", "70702000", "C", "90.00", "Facture n° FA2600001 X"),
    ]).encode("cp850")
    r = convertir(ecritures, CLIENTS)
    assert any("déséquilibrée" in a for a in r.alertes)


def test_export_excel():
    df = convertir(ECRITURES, CLIENTS).ecritures
    ws = load_workbook(io.BytesIO(vers_excel(df))).active
    assert [c.value for c in ws[1]] == COLONNES_PENNYLANE
    assert ws["C2"].value == "411003184"
    assert ws["A2"].number_format == "yyyy-mm-dd"
